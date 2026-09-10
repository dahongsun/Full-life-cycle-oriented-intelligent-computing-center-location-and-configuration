import os
import numpy as np
from numpy.linalg import inv, eigvals, cond, LinAlgError, pinv
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

# ===== 省份列表 =====
PROVINCES = ['北京','天津','河北','山西','辽宁','吉林','黑龙江','上海','江苏','浙江',
             '安徽','福建','江西','山东','河南','湖北','湖南','广东','广西','海南',
             '重庆','四川','贵州','云南','陕西','甘肃','青海','宁夏','新疆','内蒙古']

# ===== 目录路径（改成你的）=====
BASE_DIR = r"D:\BaiduNetdiskDownload\投入产出表\31省市区42部门投入产出表1997-2017年\2017年各省份投入产出表"

# ===== 文件定位：既支持 “北京.xlsx” 也支持 “01-北京.xlsx”等 =====
def locate_province_file(base_dir: str, prov: str):
    p1 = os.path.join(base_dir, f"{prov}.xlsx")
    if os.path.exists(p1): return p1
    for i in range(1, 40):
        p2 = os.path.join(base_dir, f"{i:02d}-{prov}.xlsx")
        if os.path.exists(p2): return p2
    return None

# ===== 从 A_matrix 读取 A（42×42）与行业名 =====
def read_A_and_labels(xlsx_path: str, sheet_name: str = "A_matrix"):
    wb = load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"{os.path.basename(xlsx_path)}: 找不到工作表 {sheet_name}")
    ws = wb[sheet_name]

    # 行名 A2:A43
    row_labels = [ (ws.cell(row=1+i, column=1).value or "").strip() if isinstance(ws.cell(row=1+i, column=1).value, str)
                   else ws.cell(row=1+i, column=1).value
                   for i in range(1, 1+42) ]
    # 列名 B1:AQ1
    col_labels = [ (ws.cell(row=1, column=1+j).value or "").strip() if isinstance(ws.cell(row=1, column=1+j).value, str)
                   else ws.cell(row=1, column=1+j).value
                   for j in range(1, 1+42) ]

    # A 矩阵 B2:AQ43
    A = np.zeros((42, 42), dtype=float)
    for i in range(42):
        for j in range(42):
            v = ws.cell(row=1+i, column=1+j).value  # 注意：这里是(B2) -> 行=1+i, 列=1+j
            # 更正：openpyxl 的 B2 是 row=2, col=2；我们上方取的是 (1+i,1+j)
            # 因此应加 1 偏移：
            v = ws.cell(row=1+i+1, column=1+j+1).value
            try:
                A[i, j] = float(v)
            except (TypeError, ValueError):
                A[i, j] = 0.0

    wb.close()
    return A, row_labels, col_labels

# ===== 写 L 到新表 L_inverse（含行列名）=====
def write_L_with_labels(xlsx_path: str, L: np.ndarray,
                        row_labels, col_labels, sheet_name: str = "L_inverse"):
    wb = load_workbook(xlsx_path)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    # 写列标题（B1:AQ1）
    ws.cell(row=1, column=1, value="sector")
    for j, name in enumerate(col_labels, start=2):
        ws.cell(row=1, column=j, value=name)

    # 写行标题 + 数据（从 B2 开始）
    for i, name in enumerate(row_labels, start=2):
        ws.cell(row=i, column=1, value=name)
        for j in range(42):
            ws.cell(row=i, column=2+j, value=float(L[i-2, j]))

    wb.save(xlsx_path)
    wb.close()

# ===== 计算 L，并进行基本可逆性检查 =====
def compute_leontief_inverse(A: np.ndarray, use_pinv_if_singular: bool = True):
    n = A.shape[0]
    I = np.eye(n)
    M = I - A

    # 谱半径检查
    rho = max(abs(eigvals(A)))
    # 条件数（用2-范数），越大越病态
    try:
        kappa = cond(M)
    except LinAlgError:
        kappa = np.inf

    # 求逆
    try:
        L = inv(M)
        method = "inv"
    except LinAlgError:
        if use_pinv_if_singular:
            L = pinv(M)  # 伪逆兜底（请在论文里说明处理方式）
            method = "pinv"
        else:
            raise

    return L, float(rho), float(kappa), method

# ===== 批量处理 =====
ok, fail = [], []
for prov in PROVINCES:
    path = locate_province_file(BASE_DIR, prov)
    if not path:
        fail.append((prov, "文件不存在"))
        print(f"[MISS] {prov} -> 未找到文件")
        continue

    try:
        A, rlabels, clabels = read_A_and_labels(path, sheet_name="A_matrix")
        L, rho, kappa, method = compute_leontief_inverse(A, use_pinv_if_singular=True)
        write_L_with_labels(path, L, rlabels, clabels, sheet_name="L_inverse")
        ok.append((prov, rho, kappa, method))
        print(f"[OK] {prov}: rho(A)={rho:.4f}, cond(I-A)={kappa:.2e}, method={method}")
    except Exception as e:
        fail.append((prov, str(e)))
        print(f"[FAIL] {prov} -> {e}")

print("\n==== 总结 ====")
print("成功：", ok)
print("失败：", fail)
