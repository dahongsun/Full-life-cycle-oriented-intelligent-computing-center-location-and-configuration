import os
from openpyxl import load_workbook, Workbook

# ====== 30个省份（顺序即为读取顺序）======
PROVINCES = ['北京', '天津', '河北', '山西', '辽宁', '吉林', '黑龙江', '上海', '江苏', '浙江', '安徽', '福建', '江西', '山东', '河南', '湖北', '湖南', '广东', '广西', '海南', '重庆', '四川', '贵州', '云南', '陕西', '甘肃', '青海', '宁夏', '新疆', '内蒙古']

# ====== 基本配置 ======
BASE_DIR   = r"D:\BaiduNetdiskDownload\投入产出表\31省市区42部门投入产出表1997-2017年\2017年各省份投入产出表"
SHEET_NAME = "Sheet1"
CELL_COL   = "BI"     # 取值所在列
CELL_ROW   = 51       # 取值所在行
OUT_FILE   = os.path.join(BASE_DIR, "各个省份GDP.xlsx")
CELL_ADDR  = f"{CELL_COL}{CELL_ROW}"

def read_cell_bi51(file_path: str, sheet_name: str = SHEET_NAME, cell_addr: str = CELL_ADDR):
    wb = load_workbook(file_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"找不到工作表：{sheet_name}")
    ws = wb[sheet_name]
    v = ws[cell_addr].value
    # 尝试转成 float；失败就原样返回
    try:
        return float(v)
    except (TypeError, ValueError):
        return v

def save_results_to_xlsx(results, out_path: str):
    """
    results: list of tuples [(prov, value), ...]  按原顺序
    """
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "BI51_collect"

    # 表头
    ws_out.cell(row=1, column=1, value="省份")
    ws_out.cell(row=1, column=2, value="BI51")

    # 数据
    for i, (prov, val) in enumerate(results, start=2):
        ws_out.cell(row=i, column=1, value=prov)
        ws_out.cell(row=i, column=2, value=val)

    wb_out.save(out_path)

# ========== 方案 A：文件名 = “省名.xlsx” ==========
ok, fail = [], []
results = []
for idx, prov in enumerate(PROVINCES, start=1):
    path = os.path.join(BASE_DIR, f"{prov}.xlsx")
    try:
        val = read_cell_bi51(path, sheet_name=SHEET_NAME, cell_addr=CELL_ADDR)
        results.append((prov, val))
        ok.append(prov)
        print(f"[OK] {prov} -> {CELL_ADDR} = {val}")
    except Exception as e:
        fail.append((prov, str(e)))
        print(f"[FAIL] {prov} -> {e}")

# 将成功读取到的数据写入一个总表
if results:
    save_results_to_xlsx(results, OUT_FILE)
    print(f"\n已保存汇总到：{OUT_FILE}")

print("\n====== 完成（方案A）======")
print("成功：", ok)
print("失败：", fail)

# ========== 若实际文件为 “01-北京.xlsx, 02-天津.xlsx, …” 用方案 B（把上面方案A注释掉，这段放开）==========
# ok, fail = [], []
# results = []
# for idx, prov in enumerate(PROVINCES, start=1):
#     fn = f"{idx:02d}-{prov}.xlsx"
#     path = os.path.join(BASE_DIR, fn)
#     try:
#         val = read_cell_bi51(path, sheet_name=SHEET_NAME, cell_addr=CELL_ADDR)
#         results.append((prov, val))
#         ok.append(fn)
#         print(f"[OK] {fn} -> {CELL_ADDR} = {val}")
#     except Exception as e:
#         fail.append((fn, str(e)))
#         print(f"[FAIL] {fn} -> {e}")
#
# if results:
#     save_results_to_xlsx(results, OUT_FILE)
#     print(f"\n已保存汇总到：{OUT_FILE}")
#
# print("\n====== 完成（方案B）======")
# print("成功：", ok)
# print("失败：", fail)
